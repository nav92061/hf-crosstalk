#!/usr/bin/env python3
"""Build the submission .docx from PAPER.md.

Meta-Radiology requires the manuscript in the native word-processor format, with
figures and tables placed next to the relevant text and captions directly below
the item. This produces a single-column .docx meeting those requirements.

Usage:
    python build_docx.py [PAPER.md] [out.docx]

Requires: python-docx, Pillow.
"""
import json
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image as PILImage

BODY_PT = 11
FIG_W_IN = 6.3


def _superscript(run):
    run.font.superscript = True


def add_runs(par, text):
    """Emit text with **bold**, *italic*, `code` and {^n} superscripts."""
    token = re.compile(r"(\*\*.+?\*\*|(?<!\*)\*[^*\n]+?\*(?!\*)|`[^`]+?`|\{\^-?[\d,]+\})", re.S)
    for part in token.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1]); r.font.name = "Courier New"; r.font.size = Pt(BODY_PT - 1)
        elif part.startswith("{^"):
            s = part[2:-1]
            if not s.startswith("-") and "," in s:      # AMA: collapse consecutive runs
                ns = [int(x) for x in s.split(",")]
                out, i = [], 0
                while i < len(ns):
                    j = i
                    while j + 1 < len(ns) and ns[j + 1] == ns[j] + 1:
                        j += 1
                    out.append("%d\u2013%d" % (ns[i], ns[j]) if j - i >= 2
                               else ",".join(str(n) for n in ns[i:j + 1]))
                    i = j + 1
                s = ",".join(out)
            elif s.startswith("-"):
                s = "\u2212" + s[1:]
            r = par.add_run(s); _superscript(r)
        elif part.startswith("*") and part.endswith("*"):
            r = par.add_run(part[1:-1]); r.italic = True
        else:
            par.add_run(part)


def build(src="PAPER.md", out="hf-crosstalk-manuscript.docx",
          figdir="results/figures", manifest="figure_manifest.json"):
    md = open(src).read()
    a2f = {}
    for cand in (manifest, os.path.join("results", manifest)):
        if os.path.exists(cand):
            a2f = {int(k): v for k, v in json.load(open(cand)).items()}
            break
    if not os.path.isdir(figdir):
        figdir = "figs"

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(BODY_PT)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 2.0          # double-spaced for review
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(1)
        s.top_margin = s.bottom_margin = Inches(1)

    lines = md.split("\n")
    i = 0
    while i < len(lines):
        s = lines[i].strip()
        if not s:
            i += 1
            continue

        if s.startswith("!["):
            m = re.match(r"!\[Figure (\d)\]\(([^)]*)\)", s)
            if m:
                n = int(m.group(1))
                fn = a2f.get(n) or os.path.basename(m.group(2))
                path = os.path.join(figdir, fn)
                if os.path.exists(path):
                    w, h = PILImage.open(path).size
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.line_spacing = 1.0
                    p.add_run().add_picture(path, width=Inches(min(FIG_W_IN, FIG_W_IN)))
            i += 1
            continue

        if s.startswith("|") and i + 1 < len(lines) and set(lines[i + 1].strip()) <= set("|-: "):
            rows, j = [], i
            while j < len(lines) and lines[j].strip().startswith("|"):
                cells = [c.strip() for c in lines[j].strip().strip("|").split("|")]
                if not set("".join(cells)) <= set("-: "):
                    rows.append(cells)
                j += 1
            if rows:
                t = doc.add_table(rows=len(rows), cols=len(rows[0]))
                t.style = "Table Grid"
                for r, row in enumerate(rows):
                    for c, cell in enumerate(row[:len(rows[0])]):
                        cp = t.cell(r, c).paragraphs[0]
                        cp.paragraph_format.line_spacing = 1.0
                        add_runs(cp, cell)
                        if r == 0:
                            for run in cp.runs:
                                run.bold = True
            i = j
            continue

        if s.startswith("#"):
            lvl = len(s) - len(s.lstrip("#"))
            text = s[lvl:].strip()
            if lvl == 1:
                p = doc.add_paragraph(); p.paragraph_format.line_spacing = 1.0
                r = p.add_run(text); r.bold = True; r.font.size = Pt(16)
            else:
                p = doc.add_heading(level=min(lvl, 4))
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    run.text = ""
                add_runs(p, text)
                for run in p.runs:
                    run.font.color.rgb = RGBColor(0, 0, 0)
                    run.font.name = "Times New Roman"
            i += 1
            continue

        if s.startswith(("- ", "* ")):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.line_spacing = 2.0
            add_runs(p, s[2:])
            i += 1
            continue

        para = [s]
        j = i + 1
        while j < len(lines) and lines[j].strip() and not lines[j].lstrip().startswith(
                ("#", "!", "|", "- ", "* ")):
            para.append(lines[j].strip())
            j += 1
        text = " ".join(para)
        p = doc.add_paragraph()
        if text.lstrip().startswith("**Figure"):
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(12)
        add_runs(p, text)
        i = j

    doc.save(out)
    return out


if __name__ == "__main__":
    src = sys.argv[1] if len(sys.argv) > 1 else "PAPER.md"
    out = sys.argv[2] if len(sys.argv) > 2 else "hf-crosstalk-manuscript.docx"
    print("%s -> %s" % (src, build(src, out)))
